"""Rebuild source-call and polished-isoform inventories from local source files.

Install openpyxl, download the source files named in the report README into
``--input-dir``, and run this script. The outputs intentionally preserve
distinct source call sets and do not claim mutant RNA or protein support.
"""

import argparse
from collections import Counter
import csv
import gzip
from pathlib import Path
import re

import openpyxl


ROOT = None
OUT = Path(__file__).parent

FIELDS = [
    "sample", "source", "class", "call_id", "chrom1", "pos1", "chrom2", "pos2",
    "ref", "alt", "filter", "dna_support", "published_rna_fusion", "rna_status",
]


def vcf_rows(path, sample, source, kind):
    opener = gzip.open if str(path).endswith(".gz") else open
    with opener(path, "rt") as handle:
        for line in handle:
            if line.startswith("#"):
                continue
            cols = line.rstrip("\n").split("\t")
            chrom, pos, call_id, ref, alt, _, filt, info = cols[:8]
            if kind == "SV":
                tags = dict(item.split("=", 1) for item in info.split(";") if "=" in item)
                cls = tags.get("SVTYPE", "SV")
                chrom2 = tags.get("CHR2", chrom)
                pos2 = tags.get("END", "")
                dna_support = tags.get("SUPPORT", "")
            else:
                cls = kind
                chrom2 = pos2 = dna_support = ""
            yield dict(sample=sample, source=source, **{
                "class": cls, "call_id": call_id if call_id != "." else f"{chrom}:{pos}:{ref}>{alt}",
                "chrom1": chrom, "pos1": pos, "chrom2": chrom2, "pos2": pos2,
                "ref": ref, "alt": alt, "filter": filt, "dna_support": dna_support,
                "published_rna_fusion": "", "rna_status": "not_evaluated",
            })


def hcc_sv_rows():
    sv = openpyxl.load_workbook(ROOT / "13059_2022_2816_MOESM4_ESM.xlsx", read_only=True, data_only=True).active
    fusion = openpyxl.load_workbook(ROOT / "13059_2022_2816_MOESM9_ESM.xlsx", read_only=True, data_only=True).active
    fusion_by_id = {row[0]: row[2] for row in fusion.iter_rows(min_row=2, values_only=True)
                    if isinstance(row[0], str) and row[0].startswith("chr")}
    for row in sv.iter_rows(min_row=2, values_only=True):
        if not isinstance(row[0], str) or not row[0].startswith("chr"):
            continue
        call_id, chrom1, pos1, chrom2, pos2, _, cls, replicates = row[:8]
        yield dict(sample="HCC1395", source="Talsania_2022_S4", **{
            "class": cls, "call_id": call_id, "chrom1": chrom1, "pos1": pos1,
            "chrom2": chrom2, "pos2": pos2, "ref": "", "alt": "",
            "filter": "high_confidence", "dna_support": replicates,
            "published_rna_fusion": fusion_by_id.get(call_id, ""), "rna_status": "not_evaluated",
        })


def write_calls():
    sources = [
        (ROOT / "hcc1395-seqc2-highconf-snv.vcf.gz", "HCC1395", "SEQC2_v1.2.1", "SNV"),
        (ROOT / "hcc1395-seqc2-highconf-indel.vcf.gz", "HCC1395", "SEQC2_v1.2.1", "indel"),
        (ROOT / "Truthset_SNV_COLO829.vcf.gz", "COLO829", "SMaHT_COLO829BLT50", "SNV"),
        (ROOT / "Truthset_Indel_COLO829.vcf.gz", "COLO829", "SMaHT_COLO829BLT50", "indel"),
        (ROOT / "colo829-sv-grch38.vcf", "COLO829", "PaulinPaz_2024_v2", "SV"),
        (ROOT / "colo829-sv-2022-hg38lifted.vcf", "COLO829", "ValleInclan_2022_lifted", "SV"),
    ]
    counts = Counter()
    with gzip.open(OUT / "somatic_call_inventory.tsv.gz", "wt", newline="") as handle:
        writer = csv.DictWriter(handle, FIELDS, delimiter="\t")
        writer.writeheader()
        for path, sample, source, kind in sources:
            for row in vcf_rows(path, sample, source, kind):
                writer.writerow(row)
                counts[(sample, source, row["class"])] += 1
        for row in hcc_sv_rows():
            writer.writerow(row)
            counts[(row["sample"], row["source"], row["class"])] += 1
    return counts


def write_isoforms():
    coordinates = {}
    with (ROOT / "colo829-tumor.gff").open() as handle:
        for line in handle:
            cols = line.rstrip("\n").split("\t")
            if len(cols) != 9 or cols[2] != "transcript":
                continue
            match = re.search(r'transcript_id "([^"]+)"', cols[8])
            if match:
                coordinates[match.group(1)] = (cols[0], cols[3], cols[4], cols[6])
    with (ROOT / "colo829-fl-count.txt").open(newline="") as source, gzip.open(
        OUT / "colo829_polished_isoform_counts.tsv.gz", "wt", newline=""
    ) as target:
        reader = csv.DictReader(source)
        writer = csv.writer(target, delimiter="\t")
        writer.writerow(["isoform_id", "chrom", "start", "end", "strand", "tumor_full_length_count",
                         "normal_full_length_count", "sequence_allele_status"])
        counts = Counter()
        for row in reader:
            chrom, start, end, strand = coordinates.get(row["id"], ("", "", "", ""))
            tumor, normal = int(row["COLO829T"]), int(row["COLO829BL"])
            writer.writerow([row["id"], chrom, start, end, strand, tumor, normal, "not_evaluated"])
            counts["all"] += 1
            counts["tumor_positive"] += tumor > 0
            counts["normal_positive"] += normal > 0
            counts["tumor_gff_matched"] += bool(chrom)
    return counts


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-dir", type=Path, required=True)
    args = parser.parse_args()
    ROOT = args.input_dir
    OUT.mkdir(exist_ok=True)
    print("calls", write_calls())
    print("isoforms", write_isoforms())
